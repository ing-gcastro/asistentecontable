import os
import re
import sys
import pandas as pd
import openpyxl
from datetime import datetime
from config import CARPETA_A, ARCHIVO_PROVEEDORES, REPORTE_MAESTRO_HTML
from proveedores import cargar_excel_proveedores, buscar_proveedor
from procesador_core import extraer_datos_pdf
from reporter import obtener_facturas_procesadas

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# 📍 El Excel ahora se ubica directamente en la carpeta de entrada de la red
EXCEL_VALIDACION = os.path.join(CARPETA_A, "Validacion_Lote.xlsx")

def cerrar_excel_si_abierto():
    """Fuerza el cierre de Excel en Windows si el usuario lo dejó abierto para evitar bloqueos."""
    try:
        os.system("taskkill /f /im excel.exe >nul 2>&1")
    except Exception:
        pass

def generar_excel_lote():
    print("Iniciando Preparación de Lote (Paso 1)...")
    cerrar_excel_si_abierto()
    
    if not os.path.exists(CARPETA_A):
        print(f"❌ La carpeta de entrada no existe: {CARPETA_A}")
        return

    # Omitimos el propio Excel de validación si existe en la carpeta
    archivos = [f for f in os.listdir(CARPETA_A) if f.lower().endswith(".pdf")]
    if not archivos:
        print("📂 No hay facturas pendientes en la bandeja de entrada para procesar.")
        return
        
    print("📥 Cargando bases de proveedores, sectores e historial...")
    base_prov = cargar_excel_proveedores(ARCHIVO_PROVEEDORES)
    facturas_procesadas_historico = obtener_facturas_procesadas(REPORTE_MAESTRO_HTML)
    print(f"🛡️ Historial cargado: {len(facturas_procesadas_historico)} facturas registradas previamente.")

    datos_excel = []
    vistos_en_lote = set()
    
    for archivo in archivos:
        ruta_completa = os.path.join(CARPETA_A, archivo)
        ruta_abs = os.path.abspath(ruta_completa)
        link_formula = f'=HYPERLINK("{ruta_abs}", "Ver PDF")'
        print(f"  Analizando: {archivo}")
        print(f"   └─ Link al PDF: file:///{ruta_abs.replace('\\', '/')}")
        
        exito, datos = extraer_datos_pdf(ruta_completa, remitente="")
        
        # ─── CASO 1: NO ES FACTURA O ES CATÁLOGO (>3 PÁGS) ────────────────────────
        if not exito:
            print(f"  ⚠️ Documento no válido/descartado ({datos}). Marcando para RECHAZAR.")
            datos_excel.append({
                "Archivo_Original": archivo,
                "Ver_PDF": link_formula,
                "Empresa": "TELCOM VENTURES",
                "Razón_Social": "NO FACTURA",
                "Sector": "GENERAL",
                "Tipo": "N/A",
                "Nro_Comprobante": "N/A",
                "Subtotal": "",
                "OC_Detectada": "",
                "Adjuntos_Extra": "",
                "Tarjeta": "NO",
                "Autorizar_Directo": "NO",
                "Nombre_Propuesto": "NO FACTURA",
                "Estado_Historial": "NO ES FACTURA / RECHAZAR",
                "Acción": "RECHAZAR"
            })
            continue

        # ─── CASO 2: DOCUMENTO VÁLIDO -> BUSCAR PROVEEDOR ─────────────────────────
        nombre_final, sector, es_lista_negra = buscar_proveedor(datos["cuit"], datos["razon_social"], base_prov)
        
        tipo_str = datos["tipo"] if datos["tipo"] else "FC"
        nro_str = datos["numero"] if datos["numero"] else "S_N"
        
        mes_actual = int(datetime.now().month)
        meses = {1:"ENERO", 2:"FEBRERO", 3:"MARZO", 4:"ABRIL", 5:"MAYO", 6:"JUNIO", 
                 7:"JULIO", 8:"AGOSTO", 9:"SEPTIEMBRE", 10:"OCTUBRE", 11:"NOVIEMBRE", 12:"DICIEMBRE"}
        nombre_mes = meses.get(mes_actual, "MES")
        
        sector_str = sector if sector else "GENERAL"
        
        oc_extraida_limpia = str(datos.get("oc_extraida", "")).strip()
        if oc_extraida_limpia.endswith('.0'): oc_extraida_limpia = oc_extraida_limpia[:-2]
        
        es_tarjeta = datos.get("es_tarjeta", False) or bool(re.search(r"\b(visa|visacorpo|visa corpo|tarjeta)\b", archivo, re.IGNORECASE))
        prefijo_tarjeta = "(VISA CORPO) " if es_tarjeta else ""
        es_eurosat = datos.get("es_eurosat", False)
        empresa_str = "EUROSAT" if es_eurosat else "TELCOM VENTURES"
        prefijo_empresa = "(EUROSAT) " if es_eurosat else ""
        
        nombre_sugerido = f"{prefijo_tarjeta}{prefijo_empresa}({nombre_mes}) {nombre_final} {tipo_str} {nro_str} {sector_str}"
        if oc_extraida_limpia:
            nombre_sugerido += f" -OC {oc_extraida_limpia}"

        # ─── CASO 3: DETECCIÓN DE DUPLICADAS (HISTORIAL + EN MISMO LOTE) ──────────
        razon_limpia = re.sub(r'^(\([^)]+\)\s*)+', '', nombre_final).strip().upper()
        identificador_prov = str(datos.get("cuit", "")).strip() or razon_limpia

        es_duplicada_historico = False
        if str(nro_str).strip() not in ["SinNumero", "N/A", "S_N"]:
            for (r_hist, n_hist) in facturas_procesadas_historico:
                if n_hist == str(nro_str).strip() and (r_hist in razon_limpia or razon_limpia in r_hist):
                    es_duplicada_historico = True
                    break

        clave_lote = (identificador_prov, str(nro_str).strip())
        es_duplicada_lote = False
        if str(nro_str).strip() not in ["SinNumero", "N/A", "S_N"] and clave_lote in vistos_en_lote:
            es_duplicada_lote = True
        else:
            if str(nro_str).strip() not in ["SinNumero", "N/A", "S_N"]:
                vistos_en_lote.add(clave_lote)

        if es_duplicada_historico:
            print(f"  🛑 DUPLICADA EN HISTORIAL: {nombre_final} {nro_str} ya fue procesada anteriormente.")
            datos_excel.append({
                "Archivo_Original": archivo,
                "Ver_PDF": link_formula,
                "Empresa": empresa_str,
                "Razón_Social": nombre_final,
                "Sector": sector_str,
                "Tipo": tipo_str,
                "Nro_Comprobante": str(nro_str),
                "Subtotal": datos.get("subtotal", ""),
                "OC_Detectada": oc_extraida_limpia,
                "Adjuntos_Extra": "",
                "Tarjeta": "SI" if es_tarjeta else "NO",
                "Autorizar_Directo": "NO",
                "Nombre_Propuesto": f"{nombre_sugerido} [DUPLICADA HISTORIAL]",
                "Estado_Historial": "YA PROCESADA (HISTORIAL)",
                "Acción": "RECHAZAR"
            })
        elif es_duplicada_lote:
            print(f"  🛑 DUPLICADA EN LOTE: {nombre_final} {nro_str} aparece varias veces en la misma carpeta.")
            datos_excel.append({
                "Archivo_Original": archivo,
                "Ver_PDF": link_formula,
                "Empresa": empresa_str,
                "Razón_Social": nombre_final,
                "Sector": sector_str,
                "Tipo": tipo_str,
                "Nro_Comprobante": str(nro_str),
                "Subtotal": datos.get("subtotal", ""),
                "OC_Detectada": oc_extraida_limpia,
                "Adjuntos_Extra": "",
                "Tarjeta": "SI" if es_tarjeta else "NO",
                "Autorizar_Directo": "NO",
                "Nombre_Propuesto": f"{nombre_sugerido} [DUPLICADA EN LOTE]",
                "Estado_Historial": "DUPLICADA EN MISMO LOTE",
                "Acción": "RECHAZAR"
            })
        else:
            accion_defecto = "RECHAZAR" if es_lista_negra else "APROBAR"
            datos_excel.append({
                "Archivo_Original": archivo,
                "Ver_PDF": link_formula,
                "Empresa": empresa_str,
                "Razón_Social": nombre_final,
                "Sector": sector_str,
                "Tipo": tipo_str,
                "Nro_Comprobante": str(nro_str),
                "Subtotal": datos.get("subtotal", ""),
                "OC_Detectada": oc_extraida_limpia,
                "Adjuntos_Extra": "",
                "Tarjeta": "SI" if es_tarjeta else "NO",
                "Autorizar_Directo": "NO",
                "Nombre_Propuesto": nombre_sugerido,
                "Estado_Historial": "NUEVA",
                "Acción": accion_defecto
            })
        
    df = pd.DataFrame(datos_excel)
    
    try:
        with pd.ExcelWriter(EXCEL_VALIDACION, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name="Lote_Pendiente")
            worksheet = writer.sheets['Lote_Pendiente']
            
            # Formato de celdas para hipervínculos (Azul y Subrayado)
            font_hyperlink = openpyxl.styles.Font(color="0000FF", underline="single")
            
            ver_pdf_col_idx = None
            accion_col_idx = None
            tarjeta_col_idx = None
            autorizar_col_idx = None
            
            for col_idx, cell in enumerate(worksheet[1], 1):
                col_name = str(cell.value).strip()
                if col_name == "Ver_PDF":
                    ver_pdf_col_idx = col_idx
                elif col_name == "Acción":
                    accion_col_idx = col_idx
                elif col_name == "Tarjeta":
                    tarjeta_col_idx = col_idx
                elif col_name == "Autorizar_Directo":
                    autorizar_col_idx = col_idx
            
            if ver_pdf_col_idx:
                for r_idx in range(2, worksheet.max_row + 1):
                    cell = worksheet.cell(row=r_idx, column=ver_pdf_col_idx)
                    cell.font = font_hyperlink

            # ─── AGREGAR VALIDACIÓN DE DATOS (DESPLEGABLES EN EXCEL) ───────────
            from openpyxl.worksheet.datavalidation import DataValidation
            
            if accion_col_idx:
                col_letter = openpyxl.utils.get_column_letter(accion_col_idx)
                dv_accion = DataValidation(type="list", formula1='"APROBAR,RECHAZAR,REVISION"', allow_blank=False)
                worksheet.add_data_validation(dv_accion)
                dv_accion.add(f"{col_letter}2:{col_letter}{worksheet.max_row}")

            if tarjeta_col_idx:
                col_letter = openpyxl.utils.get_column_letter(tarjeta_col_idx)
                dv_tarjeta = DataValidation(type="list", formula1='"NO,SI"', allow_blank=False)
                worksheet.add_data_validation(dv_tarjeta)
                dv_tarjeta.add(f"{col_letter}2:{col_letter}{worksheet.max_row}")

            if autorizar_col_idx:
                col_letter = openpyxl.utils.get_column_letter(autorizar_col_idx)
                dv_autorizar = DataValidation(type="list", formula1='"NO,SI"', allow_blank=False)
                worksheet.add_data_validation(dv_autorizar)
                dv_autorizar.add(f"{col_letter}2:{col_letter}{worksheet.max_row}")

            for col in worksheet.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        val_str = str(cell.value)
                        if val_str.startswith("=HYPERLINK"):
                            val_len = 10
                        else:
                            val_len = len(val_str)
                        if val_len > max_length:
                            max_length = val_len
                    except:
                        pass
                adjusted_width = (max_length + 2)
                worksheet.column_dimensions[column].width = min(max(adjusted_width, 10), 50)
                
        print(f"\n✅ ¡Lote generado con éxito en la carpeta de red!")
        print(f"👉 Excel generado en: '{EXCEL_VALIDACION}'")
        print(f"👉 Opciones de Acción disponibles en Excel: [APROBAR, RECHAZAR, REVISION]")
        print(f"👉 Revisa las Acciones, edita Sector u OC si hace falta, guarda y cierra el Excel.")
        print(f"👉 Luego, ejecuta el Paso 2 (Aplicar Lote).")
    except Exception as e:
        print(f"❌ Error al guardar el Excel: {e}")

if __name__ == "__main__":
    generar_excel_lote()
    input("\nPresiona ENTER para salir...")
