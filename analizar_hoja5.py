import openpyxl
import re
from unicodedata import normalize as uni_norm

def clean_str(text):
    if not text:
        return ""
    text = uni_norm('NFD', str(text)).encode('ascii', 'ignore').decode('utf-8')
    text = text.upper().strip()
    text = re.sub(r'\b(S\.?R\.?L\.?|S\.?A\.?|S\.?H\.?|S\.?A\.?C\.?I\.?F\.?I\.?|S\.?A\.?I\.?C\.?|S\.?A\.?U\.?|S\.?A\.?S\.?|INC|LLC)\b', '', text)
    text = re.sub(r'[^A-Z0-9\s]', ' ', text)
    text = ' '.join(text.split())
    return text

wb = openpyxl.load_workbook(r'C:\Users\gcastro\ROBOTS\data\proveedores (2).xls.xlsx', data_only=True)
ws5 = wb['Hoja5']
ws1 = wb['Hoja1']

# Target data (Hoja1)
hoja1_rows = []
for r in range(2, ws1.max_row + 1):
    c_num = ws1.cell(r, 1).value
    desc = ws1.cell(r, 2).value  # Col B (Descripcion / Razon Social)
    cuit = ws1.cell(r, 6).value
    cur_email = ws1.cell(r, 12).value  # Col L (EMail)
    if desc:
        desc_str = str(desc).strip()
        hoja1_rows.append({
            'row': r,
            'cocuenta': c_num,
            'desc_raw': desc_str,
            'desc_clean': clean_str(desc_str),
            'cuit': str(cuit).strip() if cuit else '',
            'cur_email': str(cur_email).strip() if cur_email else ''
        })

# Source data (Hoja5)
# In Hoja5:
# Col D (4): Razon Social / Contact
# Col E (5): Direccion de Mail
# Col K (11): Extracted provider name from subject
hoja5_map = {} # provider_clean -> {raw_names: set(), emails: set(), rows: list()}

for r in range(2, ws5.max_row + 1):
    col_d = ws5.cell(r, 4).value
    col_e = ws5.cell(r, 5).value
    col_k = ws5.cell(r, 11).value
    
    email = str(col_e).strip() if col_e else ""
    if not email or email.lower() == 'none':
        continue
        
    names_to_try = []
    if col_k and str(col_k).strip():
        names_to_try.append(str(col_k).strip())
    if col_d and str(col_d).strip():
        names_to_try.append(str(col_d).strip())
        
    for raw_n in names_to_try:
        clean_n = clean_str(raw_n)
        # Skip generic internal names / emails used as name
        if not clean_n or clean_n in ['COMPRAS', 'PAGO A PROVEEDORES', 'LISTA NEGRA'] or '@' in raw_n:
            continue
            
        if clean_n not in hoja5_map:
            hoja5_map[clean_n] = {
                'raw_names': set(),
                'emails': set(),
                'rows': []
            }
        hoja5_map[clean_n]['raw_names'].add(raw_n)
        hoja5_map[clean_n]['emails'].add(email)
        hoja5_map[clean_n]['rows'].append(r)

print(f"Total unique valid providers in Hoja5: {len(hoja5_map)}")
print(f"Total rows in Hoja1: {len(hoja1_rows)}")

matched_hoja5 = {}
unmatched_hoja5 = {}

for clean_n, data in hoja5_map.items():
    # Try to find match in Hoja1
    # 1. Clean match
    matches = [h for h in hoja1_rows if h['desc_clean'] == clean_n]
    
    # 2. Substring match
    if not matches and len(clean_n) >= 4:
        for h in hoja1_rows:
            h_c = h['desc_clean']
            if not h_c: continue
            if clean_n in h_c or h_c in clean_n:
                ratio = len(clean_n) / len(h_c) if len(h_c) > 0 else 0
                if 0.4 < ratio < 2.5:
                    matches.append(h)
                    
    if matches:
        matched_hoja5[clean_n] = {
            'hoja5_data': data,
            'targets': matches
        }
    else:
        unmatched_hoja5[clean_n] = data

print(f"\nMatched providers: {len(matched_hoja5)}")
print(f"Unmatched providers: {len(unmatched_hoja5)}")

print("\n--- UNMATCHED PROVIDERS FROM HOJA5 ---")
for clean_n, data in unmatched_hoja5.items():
    raws = " / ".join(data['raw_names'])
    emails = ", ".join(data['emails'])
    print(f"Hoja5: [{raws}] (Mails: {emails})")
