import openpyxl
import re
from unicodedata import normalize as uni_norm

def clean_str(text):
    if not text:
        return ""
    # Remove accents
    text = uni_norm('NFD', str(text)).encode('ascii', 'ignore').decode('utf-8')
    text = text.upper().strip()
    # Remove common corporate designations for comparison
    text = re.sub(r'\b(S\.?R\.?L\.?|S\.?A\.?|S\.?H\.?|S\.?A\.?C\.?I\.?F\.?I\.?|S\.?A\.?I\.?C\.?)\b', '', text)
    # Remove punctuation
    text = re.sub(r'[^A-Z0-9\s]', ' ', text)
    # Remove extra spaces
    text = ' '.join(text.split())
    return text

wb_prov = openpyxl.load_workbook(r'C:\Users\gcastro\ROBOTS\data\proveedores (2).xls.xlsx', data_only=True)
ws1 = wb_prov['Hoja1']

# Target data (Hoja1)
hoja1_rows = []
for r in range(2, ws1.max_row + 1):
    c_num = ws1.cell(r, 1).value  # CoCuenta
    rs = ws1.cell(r, 2).value     # Descripcion (Razon Social)
    cuit = ws1.cell(r, 6).value   # Numero_Documento
    email = ws1.cell(r, 12).value # EMail
    if rs:
        hoja1_rows.append({
            'row': r,
            'cocuenta': c_num,
            'rs': str(rs).strip(),
            'rs_clean': clean_str(rs),
            'cuit': str(cuit).strip() if cuit else '',
            'email': str(email).strip() if email else ''
        })

wb_nom = openpyxl.load_workbook(r'C:\Users\gcastro\ROBOTS\data\Nombreproveedores.xlsx', data_only=True)
ws_source = wb_nom['Hoja2']

source_list = []
for r in range(2, ws_source.max_row + 1):
    email = ws_source.cell(r, 1).value # Correo Proveedor
    rs = ws_source.cell(r, 2).value    # RAZON SOCIAL
    sector = ws_source.cell(r, 3).value
    if rs or email:
        rs_str = str(rs).strip() if rs else ''
        source_list.append({
            'row': r,
            'email': str(email).strip() if email else '',
            'rs': rs_str,
            'rs_clean': clean_str(rs_str),
            'sector': str(sector).strip() if sector else ''
        })

print(f"Total source rows in Nombreproveedores.xlsx: {len(source_list)}")
print(f"Total rows in Hoja1 of proveedores (2).xls.xlsx: {len(hoja1_rows)}")

matched_count = 0
matched_pairs = []
unmatched_source = []

for src in source_list:
    src_rs = src['rs']
    src_clean = src['rs_clean']
    
    if not src_rs or src_rs.upper() == 'LISTA NEGRA':
        unmatched_source.append({'source': src, 'reason': 'Ignorado (LISTA NEGRA / Sin Nombre)'})
        continue
    
    # 1. Exact match on raw string
    exact_m = [h for h in hoja1_rows if h['rs'].upper() == src_rs.upper()]
    
    # 2. Clean match
    clean_m = [h for h in hoja1_rows if h['rs_clean'] == src_clean and src_clean != ''] if not exact_m else []
    
    # 3. Substring / Prefix match
    sub_m = []
    if not exact_m and not clean_m and len(src_clean) >= 3:
        for h in hoja1_rows:
            h_clean = h['rs_clean']
            if not h_clean: continue
            if src_clean in h_clean or h_clean in src_clean:
                # check similarity
                ratio = len(src_clean) / len(h_clean) if len(h_clean) > 0 else 0
                if ratio > 0.4 and ratio < 2.5:
                    sub_m.append(h)
                    
    candidates = exact_m or clean_m or sub_m
    
    if candidates:
        matched_count += 1
        match_type = 'EXACTO' if exact_m else ('LIMPIO' if clean_m else 'PARCIAL')
        matched_pairs.append({'source': src, 'candidates': candidates, 'type': match_type})
    else:
        unmatched_source.append({'source': src, 'reason': 'No encontrado en Hoja1'})

print(f"\n--- MATCH SUMMARY ---")
print(f"Matched source providers: {matched_count} / {len(source_list)}")
print(f"Unmatched source providers: {len(unmatched_source)} / {len(source_list)}")

print("\n--- UNMATCHED PROVIDERS DETAILS ---")
for item in unmatched_source:
    src = item['source']
    print(f"Fila {src['row']}: Razón Social: '{src['rs']}' | Email: '{src['email']}' | Motivo: {item['reason']}")

print("\n--- SAMPLE MATCHES (First 10) ---")
for p in matched_pairs[:10]:
    src = p['source']
    cand = p['candidates'][0]
    print(f"Source: [{src['rs']}] -> Hoja1 (Fila {cand['row']}): [{cand['rs']}] | Email: {src['email']} (Match: {p['type']})")
